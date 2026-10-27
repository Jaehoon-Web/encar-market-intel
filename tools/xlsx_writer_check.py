# -*- coding: utf-8 -*-
"""
xlsx_writer_check.py — 사이트 엑셀 작성기(web/lib/xlsx_writer.js)가 정상 .xlsx 를 만드는지 확인

    python tools/xlsx_writer_check.py        (Node.js · openpyxl 필요)

node 로 시험 파일을 만들고 openpyxl 로 열어 시트 이름·값·서식(굵게·천 단위 쉼표·열 너비·첫 행 고정)을 비교한다. 다르면 종료 코드 1
"""
import os, sys, zipfile, tempfile, subprocess
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    out = os.path.join(tempfile.gettempdir(), "xlsx_writer_check.xlsx")
    subprocess.check_call(["node", os.path.join(HERE, "xlsx_writer_check.js"), out])
    ok = zipfile.ZipFile(out).testzip() is None
    wb = openpyxl.load_workbook(out)
    a, b = wb[wb.sheetnames[0]], wb[wb.sheetnames[1]]
    checks = {
        "zip": ok,
        "시트 이름": wb.sheetnames == ["산정 결과", "비교 매물 목록 "],
        "값": [r for r in a.iter_rows(values_only=True)][1:] == [
            ("추정 시세(만원)", 2149), ("대상 차", None), ("모델", "현대 더 뉴 그랜저 IG"), ("특수문자", '<a & "b"> \'c\''),
            ("소수", 17.25), ("음수", -30), ("줄바꿈", "첫 줄\n둘째 줄")],
        "머리글 굵게": a["A1"].font.b is True,
        "천 단위 쉼표": a["B2"].number_format == "#,##0" and b["C2"].number_format == "#,##0",
        "열 너비": a.column_dimensions["A"].width == 24,
        "첫 행 고정": b.freeze_panes == "A2",
        "행 수": b.max_row == 31 and b["B31"].value == 42500029,
    }
    for k, v in checks.items():
        print(("✓ " if v else "✗ ") + k)
    sys.exit(0 if all(checks.values()) else 1)


if __name__ == "__main__":
    main()
